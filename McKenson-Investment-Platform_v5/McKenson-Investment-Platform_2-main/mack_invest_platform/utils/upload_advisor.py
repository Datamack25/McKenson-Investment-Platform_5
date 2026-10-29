# utils/upload_advisor.py  —  MAM v5.5 · Conseil du comité sur un portefeuille importé (CSV / Excel)
"""
Moteur (sans interface) de l'onglet « 📥 Conseil portefeuille » :

1. lecture et contrôle d'un fichier CSV ou Excel (format décrit dans FORMAT_ROWS) ;
2. valorisation au dernier cours (ou au montant indiqué), secteur, classe, région ;
3. avis du comité MAM, réallocation et optimisation (utils.committee) — appelés par la page ;
4. scénarios de marché (stress tests du desk + scénario personnalisé) ;
5. classement des secteurs et des régions qui performent le mieux, en tenant compte du contexte ;
6. exports : Excel (avec graphiques), CSV, image PNG — chacun porte l'avertissement légal.

Aucun ordre n'est passé : c'est une analyse pédagogique. Les montants sont supposés exprimés
dans une seule devise (dollar US par défaut, comme le reste de la plateforme).
"""
from __future__ import annotations

import io
import math
import re
import unicodedata
from datetime import datetime

import numpy as np
import pandas as pd

from utils import advisor as A
from utils.legal import AUTHOR, CREDIT, DISCLAIMER

UPLOAD_VERSION = "1.0"
MAX_LINES = 200
MAX_BYTES = 5 * 1024 * 1024

# ── Avertissement imprimé après chaque analyse et dans chaque fichier ────────────
WARNING = (f"CE N'EST PAS UN CONSEIL FINANCIER. Analyse automatique à visée éducative, produite par des membres de comité "
           f"virtuels à partir de données publiques ; elle ne constitue ni un conseil en investissement, ni une "
           f"recommandation personnalisée. {AUTHOR} n'est pas responsable des décisions prises à partir de cette "
           f"analyse et ne s'engage pas sur ses résultats. Les performances passées ne préjugent pas des "
           f"performances futures ; tout investissement comporte un risque de perte en capital.")

# ══════════════════════════════════════════════════════════════════════════════
#  1. FORMAT DU FICHIER
# ══════════════════════════════════════════════════════════════════════════════
COLUMNS = {
    "ticker":     ["ticker", "symbole", "symbol", "code", "code_yahoo", "valeur_ticker"],
    "quantite":   ["quantite", "qty", "quantity", "nombre", "nb", "shares", "parts", "titres"],
    "montant":    ["montant", "valeur", "value", "amount", "market_value", "valeur_de_marche", "valorisation"],
    "prix_achat": ["prix_achat", "prix_d_achat", "pru", "prix_moyen", "cost", "avg_price", "cout_moyen"],
    "nom":        ["nom", "name", "libelle", "intitule", "description"],
    "devise":     ["devise", "currency", "ccy"],
}
CASH_TICKERS = {"CASH", "LIQUIDITES", "LIQUIDITE", "ESPECES", "CASH_USD", "CASH_EUR", "MONETAIRE"}

FORMAT_ROWS = [
    ("ticker", "Obligatoire", "Symbole Yahoo Finance de la ligne : AAPL, MSFT, SPY, CW8.PA, MC.PA, BTC-USD, GC=F… "
               "Une ligne CASH (ou LIQUIDITES) indique les liquidités.", "AAPL"),
    ("quantite", "Obligatoire si pas de montant", "Nombre de titres ou de parts détenus (nombre positif, décimales "
                                                  "acceptées).", "25"),
    ("montant", "Obligatoire si pas de quantité", "Valeur de la ligne dans la devise du portefeuille. Si quantité et "
                                                  "montant sont fournis, le dernier cours est utilisé et le montant sert "
                                                  "de repli.", "4 500"),
    ("prix_achat", "Facultatif", "Prix de revient unitaire (PRU) : permet de calculer la plus ou moins-value.", "150,20"),
    ("nom", "Facultatif", "Libellé libre de la ligne.", "Apple Inc."),
    ("devise", "Facultatif", "Devise d'affichage (une seule devise par fichier).", "USD"),
]
FORMAT_RULES = [
    "Fichier .csv (séparateur « ; » ou « , », encodage UTF-8 ou Latin-1) ou .xlsx (première feuille, ou feuille "
    "nommée « Portefeuille »).",
    "Une ligne par position ; la première ligne contient les noms de colonnes (majuscules et accents indifférents ; "
    "les synonymes Symbol, Quantity, Valeur, PRU… sont reconnus).",
    "Nombres au format français (1 234,56) ou anglais (1,234.56) ; les symboles € et $ sont ignorés.",
    f"Au plus {MAX_LINES} lignes et 5 Mo. Les lignes en double sont additionnées.",
    "Pas de positions vendeuses (quantités négatives) ni de produits dérivés : l'analyse porte sur un portefeuille "
    "au comptant.",
]


def template_frame() -> pd.DataFrame:
    return pd.DataFrame([
        dict(ticker="SPY", quantite=40, montant=None, prix_achat=480.0, nom="S&P 500 ETF", devise="USD"),
        dict(ticker="AAPL", quantite=60, montant=None, prix_achat=175.5, nom="Apple", devise="USD"),
        dict(ticker="CW8.PA", quantite=None, montant=12000, prix_achat=None, nom="Amundi MSCI World", devise="USD"),
        dict(ticker="TLT", quantite=80, montant=None, prix_achat=95.0, nom="Obligations US 20 ans +", devise="USD"),
        dict(ticker="GLD", quantite=15, montant=None, prix_achat=200.0, nom="Or", devise="USD"),
        dict(ticker="BTC-USD", quantite=0.05, montant=None, prix_achat=55000.0, nom="Bitcoin", devise="USD"),
        dict(ticker="CASH", quantite=None, montant=8000, prix_achat=None, nom="Liquidités", devise="USD"),
    ])


def template_csv() -> bytes:
    return template_frame().to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")


def template_xlsx() -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        template_frame().to_excel(xw, sheet_name="Portefeuille", index=False)
        pd.DataFrame(FORMAT_ROWS, columns=["Colonne", "Statut", "Contenu", "Exemple"]).to_excel(
            xw, sheet_name="Format", index=False)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
#  2. LECTURE ET CONTRÔLE
# ══════════════════════════════════════════════════════════════════════════════
def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s).strip().lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def _num(v) -> float:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return float("nan")
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("\xa0", "").replace(" ", "")
    s = re.sub(r"[€$£%]|EUR|USD", "", s, flags=re.I)
    if not s:
        return float("nan")
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return float("nan")


def _read_any(data: bytes, filename: str) -> pd.DataFrame:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        xl = pd.ExcelFile(io.BytesIO(data))
        sheet = next((s for s in xl.sheet_names if _norm(s) == "portefeuille"), xl.sheet_names[0])
        return xl.parse(sheet, dtype=object)
    for enc in ("utf-8-sig", "latin-1"):
        try:
            txt = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sep = ";" if txt.splitlines()[0].count(";") >= txt.splitlines()[0].count(",") else ","
    return pd.read_csv(io.StringIO(txt), sep=sep, dtype=object)


def parse_file(data: bytes, filename: str) -> dict:
    """Renvoie dict(ok, lignes: DataFrame[ticker, quantite, montant, prix_achat, nom, devise], cash, erreurs, remarques)."""
    errors, notes = [], []
    if not data:
        return dict(ok=False, erreurs=["Fichier vide."], remarques=[])
    if len(data) > MAX_BYTES:
        return dict(ok=False, erreurs=["Fichier trop volumineux (5 Mo au maximum)."], remarques=[])
    try:
        raw = _read_any(data, filename)
    except Exception as e:
        return dict(ok=False, erreurs=[f"Lecture impossible ({type(e).__name__}) : vérifiez qu'il s'agit d'un "
                                       "fichier .csv ou .xlsx conforme au modèle."], remarques=[])
    mapping = {}
    for col in raw.columns:
        n = _norm(col)
        for key, syn in COLUMNS.items():
            if n in syn and key not in mapping.values():
                mapping[col] = key
                break
    df = raw.rename(columns=mapping)
    missing_cols = [c for c in raw.columns if c not in mapping]
    if missing_cols:
        notes.append("Colonnes ignorées : " + ", ".join(map(str, missing_cols)) + ".")
    if "ticker" not in df.columns:
        return dict(ok=False, erreurs=["Colonne « ticker » introuvable (synonymes acceptés : symbole, symbol, code)."],
                    remarques=notes)
    if "quantite" not in df.columns and "montant" not in df.columns:
        return dict(ok=False, erreurs=["Il faut une colonne « quantite » ou une colonne « montant »."], remarques=notes)
    for c in ("quantite", "montant", "prix_achat", "nom", "devise"):
        if c not in df.columns:
            df[c] = None
    df = df[["ticker", "quantite", "montant", "prix_achat", "nom", "devise"]].copy()
    df["ticker"] = df["ticker"].astype(str).str.strip().str.upper().replace({"NAN": "", "NONE": ""})
    df = df[df["ticker"] != ""]
    if len(df) > MAX_LINES:
        return dict(ok=False, erreurs=[f"{len(df)} lignes : {MAX_LINES} au maximum."], remarques=notes)
    for c in ("quantite", "montant", "prix_achat"):
        df[c] = df[c].map(_num)
    cash = 0.0
    keep = []
    for i, r in df.iterrows():
        line = i + 2                                     # numéro de ligne dans le fichier (en-tête = 1)
        t = r["ticker"]
        if _norm(t).upper() in CASH_TICKERS:
            amt = r["montant"] if np.isfinite(r["montant"]) else r["quantite"]
            if np.isfinite(amt) and amt >= 0:
                cash += float(amt)
            else:
                errors.append(f"Ligne {line} : montant de liquidités invalide.")
            continue
        if not re.fullmatch(r"[A-Z0-9^.\-=_&]{1,20}", t):
            errors.append(f"Ligne {line} : ticker « {t} » invalide.")
            continue
        q, m = r["quantite"], r["montant"]
        if (np.isfinite(q) and q < 0) or (np.isfinite(m) and m < 0):
            errors.append(f"Ligne {line} ({t}) : quantité ou montant négatif (positions vendeuses non prises en charge).")
            continue
        if not ((np.isfinite(q) and q > 0) or (np.isfinite(m) and m > 0)):
            errors.append(f"Ligne {line} ({t}) : ni quantité ni montant positif.")
            continue
        keep.append(r)
    lines = pd.DataFrame(keep, columns=df.columns)
    if not lines.empty:
        dup = lines["ticker"].duplicated(keep=False)
        if dup.any():
            notes.append("Lignes en double additionnées : " + ", ".join(sorted(set(lines.loc[dup, "ticker"]))) + ".")
        lines = lines.groupby("ticker", as_index=False).agg(
            quantite=("quantite", lambda s: s.sum(min_count=1)), montant=("montant", lambda s: s.sum(min_count=1)),
            prix_achat=("prix_achat", "first"), nom=("nom", "first"), devise=("devise", "first"))
    devs = {str(d).upper() for d in df["devise"].dropna() if str(d).strip()}
    if len(devs) > 1:
        notes.append("Plusieurs devises indiquées (" + ", ".join(sorted(devs)) + ") : les montants sont additionnés "
                     "sans conversion. Convertissez-les dans une seule devise pour un résultat exact.")
    ok = not lines.empty or cash > 0
    if not ok and not errors:
        errors.append("Aucune ligne exploitable.")
    return dict(ok=ok and not lines.empty, lignes=lines, cash=cash, erreurs=errors, remarques=notes,
                devise=(sorted(devs)[0] if len(devs) == 1 else "USD"))


# ══════════════════════════════════════════════════════════════════════════════
#  3. VALORISATION ET DESCRIPTION
# ══════════════════════════════════════════════════════════════════════════════
_EU = (".PA", ".DE", ".AS", ".MI", ".MC", ".L", ".SW", ".BR", ".ST", ".CO", ".HE", ".OL", ".LS", ".IR", ".VI")
_ASIA = (".HK", ".T", ".SS", ".SZ", ".KS", ".TW", ".NS", ".BO", ".SI", ".AX", ".JK")
_LATAM = (".SA", ".MX")


def region_of(ticker: str) -> str:
    t = ticker.upper()
    try:
        from utils.etf_universe import etf_meta
        m = etf_meta(t)
        if m and m.get("zone"):
            return str(m["zone"])
    except Exception:
        pass
    if t.endswith(("-USD", "-EUR")):
        return "Mondial (crypto)"
    if t.endswith("=F"):
        return "Mondial (contrat à terme)"
    if t.endswith("=X"):
        return "Devises"
    if t.endswith(_EU):
        return "Europe"
    if t.endswith(_ASIA):
        return "Asie-Pacifique"
    if t.endswith(_LATAM):
        return "Amérique latine"
    cls = A.classify_ticker(t) or "Actions US"
    return {"Actions Europe": "Europe", "Actions Asie & émergents": "Asie & émergents", "Or": "Mondial (or)",
            "Matières premières": "Mondial (matières premières)", "Crypto": "Mondial (crypto)"}.get(cls, "États-Unis")


def value_lines(parsed: dict, live: dict, sources: dict | None = None) -> tuple[pd.DataFrame, list[str]]:
    """Tableau enrichi des positions + remarques de valorisation."""
    from utils.sectors import sector_of
    sources = sources or {}
    rows, notes = [], []
    for _, r in parsed["lignes"].iterrows():
        t = r["ticker"]
        p = float((live.get(t) or (0.0, 0.0))[0] or 0.0)
        src = sources.get(t, "live" if p > 0 else "indisponible")
        q, m = r["quantite"], r["montant"]
        if np.isfinite(q) and q > 0 and p > 0 and src == "live":
            val, base = q * p, "dernier cours"
        elif np.isfinite(m) and m > 0:
            val, base = m, "montant du fichier"
            if np.isfinite(q) and q > 0:
                p = m / q
        elif np.isfinite(q) and q > 0 and p > 0:
            val, base = q * p, "cours de repli"
            notes.append(f"{t} : cours en direct indisponible, cours de repli utilisé.")
        else:
            notes.append(f"{t} : aucun cours disponible et pas de montant — ligne ignorée (ajoutez la colonne montant).")
            continue
        pru = r["prix_achat"]
        pnl = (p - pru) * q if np.isfinite(pru) and np.isfinite(q) and q > 0 and p > 0 else np.nan
        rows.append(dict(Ticker=t, Nom=(r["nom"] if isinstance(r["nom"], str) and r["nom"] else ""),
                         Secteur=sector_of(t), Classe=A.classify_ticker(t) or "Actions US", Région=region_of(t),
                         Quantité=q if np.isfinite(q) else np.nan, Cours=p if p > 0 else np.nan, Valeur=float(val),
                         **{"Base de valorisation": base, "Plus-value": pnl,
                            "Plus-value %": (p / pru - 1) if np.isfinite(pru) and pru > 0 and p > 0 else np.nan}))
    df = pd.DataFrame(rows)
    if df.empty:
        return df, notes
    nav = df["Valeur"].sum() + parsed.get("cash", 0.0)
    df["Poids"] = df["Valeur"] / nav if nav else 0.0
    return df.sort_values("Valeur", ascending=False).reset_index(drop=True), notes


# ══════════════════════════════════════════════════════════════════════════════
#  4. SECTEURS ET RÉGIONS QUI PERFORMENT LE MIEUX (performance + contexte)
# ══════════════════════════════════════════════════════════════════════════════
SECTOR_ETF = {"Technologie": "XLK", "Santé": "XLV", "Finance": "XLF", "Énergie": "XLE", "Services publics": "XLU",
              "Industrie": "XLI", "Consommation de base": "XLP", "Consommation discrétionnaire": "XLY",
              "Matériaux": "XLB", "Immobilier": "XLRE", "Communication": "XLC"}
# Rotation sectorielle classique selon la phase du cycle (+1 favorisé, −1 pénalisé)
SECTOR_REGIME = {
    "Expansion désinflationniste": {"Technologie": 1, "Consommation discrétionnaire": 1, "Industrie": 1, "Communication": 1,
                                    "Finance": 0.5, "Services publics": -1, "Consommation de base": -0.5, "Énergie": -0.5},
    "Surchauffe / reflation": {"Énergie": 1, "Matériaux": 1, "Industrie": 1, "Finance": 0.5, "Immobilier": -0.5,
                               "Services publics": -1, "Technologie": -0.5, "Consommation de base": -0.5},
    "Ralentissement désinflationniste": {"Santé": 1, "Consommation de base": 1, "Services publics": 1, "Immobilier": 0.5,
                                         "Technologie": 0.5, "Énergie": -1, "Matériaux": -1, "Industrie": -0.5},
    "Stagflation": {"Énergie": 1, "Santé": 1, "Consommation de base": 1, "Services publics": 0.5, "Technologie": -1,
                    "Consommation discrétionnaire": -1, "Immobilier": -1, "Finance": -0.5},
    "Aversion au risque": {"Consommation de base": 1, "Santé": 1, "Services publics": 1, "Finance": -1,
                           "Consommation discrétionnaire": -1, "Technologie": -0.5, "Énergie": -0.5, "Matériaux": -0.5},
}
REGIONS = [  # (pays / zone, ETF, code ISO-3 pour la carte, grande région, classe du desk)
    ("États-Unis", "SPY", "USA", "Amérique du Nord", "Actions US"),
    ("Canada", "EWC", "CAN", "Amérique du Nord", "Actions US"),
    ("Mexique", "EWW", "MEX", "Amérique latine", "Actions Asie & émergents"),
    ("Brésil", "EWZ", "BRA", "Amérique latine", "Actions Asie & émergents"),
    ("Royaume-Uni", "EWU", "GBR", "Europe", "Actions Europe"),
    ("France", "EWQ", "FRA", "Europe", "Actions Europe"),
    ("Allemagne", "EWG", "DEU", "Europe", "Actions Europe"),
    ("Italie", "EWI", "ITA", "Europe", "Actions Europe"),
    ("Espagne", "EWP", "ESP", "Europe", "Actions Europe"),
    ("Suisse", "EWL", "CHE", "Europe", "Actions Europe"),
    ("Japon", "EWJ", "JPN", "Asie-Pacifique", "Actions Asie & émergents"),
    ("Chine", "MCHI", "CHN", "Asie émergente", "Actions Asie & émergents"),
    ("Inde", "INDA", "IND", "Asie émergente", "Actions Asie & émergents"),
    ("Corée du Sud", "EWY", "KOR", "Asie-Pacifique", "Actions Asie & émergents"),
    ("Taïwan", "EWT", "TWN", "Asie-Pacifique", "Actions Asie & émergents"),
    ("Australie", "EWA", "AUS", "Asie-Pacifique", "Actions Asie & émergents"),
    ("Afrique du Sud", "EZA", "ZAF", "Afrique", "Actions Asie & émergents"),
]


def board_tickers() -> tuple:
    return tuple(sorted(set(SECTOR_ETF.values()) | {r[1] for r in REGIONS}))


def _perf(px: pd.DataFrame | None, t: str, days: int) -> float:
    if px is None or px.empty or t not in px.columns:
        return float("nan")
    s = px[t].dropna()
    if len(s) <= days:
        return float("nan")
    return float(s.iloc[-1] / s.iloc[-1 - days] - 1)


def _rank01(s: pd.Series) -> pd.Series:
    s = s.astype(float)
    if s.notna().sum() < 2:
        return pd.Series(0.5, index=s.index)
    return s.rank(pct=True).fillna(0.5)


def sector_board(px: pd.DataFrame | None, macro) -> pd.DataFrame:
    tilt = SECTOR_REGIME.get(macro.regime, {})
    rows = []
    for sec, etf in SECTOR_ETF.items():
        rows.append(dict(Secteur=sec, ETF=etf, **{"Perf. 1 mois": _perf(px, etf, 21), "Perf. 3 mois": _perf(px, etf, 63),
                                                   "Perf. 12 mois": _perf(px, etf, 252)}, Contexte=float(tilt.get(sec, 0))))
    df = pd.DataFrame(rows).set_index("Secteur")
    score = 0.40 * _rank01(df["Perf. 3 mois"]) + 0.20 * _rank01(df["Perf. 12 mois"]) + 0.40 * (df["Contexte"] + 1) / 2
    df["Score"] = (score * 100).round(0)
    df["Lecture du contexte"] = df["Contexte"].map(lambda v: "favorisé" if v > 0.25 else ("pénalisé" if v < -0.25 else "neutre"))
    return df.sort_values("Score", ascending=False)


def region_board(px: pd.DataFrame | None, macro) -> pd.DataFrame:
    sc = A.class_scores(macro.factors)
    rows = []
    for name, etf, iso, zone, cls in REGIONS:
        rows.append(dict(Région=name, ETF=etf, ISO3=iso, Zone=zone, **{"Perf. 1 mois": _perf(px, etf, 21),
                                                                      "Perf. 3 mois": _perf(px, etf, 63),
                                                                      "Perf. 12 mois": _perf(px, etf, 252)},
                         Contexte=float(np.tanh(sc.get(cls, 0.0) / 3)), Classe=cls))
    df = pd.DataFrame(rows).set_index("Région")
    score = 0.40 * _rank01(df["Perf. 3 mois"]) + 0.20 * _rank01(df["Perf. 12 mois"]) + 0.40 * (df["Contexte"] + 1) / 2
    df["Score"] = (score * 100).round(0)
    df["Lecture du contexte"] = df["Contexte"].map(lambda v: "favorisé" if v > 0.1 else ("pénalisé" if v < -0.1 else "neutre"))
    return df.sort_values("Score", ascending=False)


# ══════════════════════════════════════════════════════════════════════════════
#  5. SCÉNARIOS
# ══════════════════════════════════════════════════════════════════════════════
CUSTOM_GROUPS = {
    "Actions": ["Actions US", "Actions Europe", "Actions Asie & émergents"],
    "Obligations": ["Oblig. long terme", "Oblig. moyen terme", "Crédit IG", "High yield"],
    "Or": ["Or"], "Matières premières": ["Matières premières"], "Crypto": ["Crypto"],
}


def class_weights(values: dict, cash: float) -> pd.Series:
    cls_w, _ = A.portfolio_to_classes({**{t: v for t, v in values.items() if v > 0}, "CASH": cash})
    return cls_w.reindex(A.CLASSES).fillna(0.0)


def scenarios(w_before: pd.Series, w_after: pd.Series, custom: dict | None = None, nav: float = 0.0) -> pd.DataFrame:
    before, after = A.stress_test(w_before), A.stress_test(w_after)
    df = pd.DataFrame({"Portefeuille actuel": before, "Après réallocation": after})
    if custom:
        shock = pd.Series(0.0, index=A.CLASSES)
        for g, v in custom.items():
            for c in CUSTOM_GROUPS.get(g, []):
                shock[c] = v
        df.loc["Scénario personnalisé"] = [float((w_before * shock).sum()), float((w_after * shock).sum())]
    df["Écart"] = df["Après réallocation"] - df["Portefeuille actuel"]
    if nav:
        df["Perte / gain actuel ($)"] = df["Portefeuille actuel"] * nav
    return df.sort_values("Portefeuille actuel")


# ══════════════════════════════════════════════════════════════════════════════
#  6. EXPORTS
# ══════════════════════════════════════════════════════════════════════════════
def _avis_frame(rev: dict) -> pd.DataFrame:
    vote = {"ok": "Favorable", "reserve": "Favorable avec réserves", "veto": "Défavorable (veto)"}
    return pd.DataFrame([{"Membre": a["nom"], "Rôle": a["role"], "Avis": vote.get(a["vote"], a["vote"]),
                          "Titre": a["titre"], "Analyse": " • ".join(re.sub(r"\*\*", "", p) for p in a["points"]),
                          "Actions proposées": " • ".join(a["actions"])} for a in rev["avis"]])


def _realloc_frame(res: dict) -> pd.DataFrame:
    opt = res["opt"]
    cur = pd.Series(opt["actuel_lignes"], dtype=float)
    tgt = pd.Series(opt["cible_lignes"], dtype=float)
    idx = cur.index.union(tgt.index)
    from utils.sectors import sector_of
    df = pd.DataFrame({"Secteur": [sector_of(t) for t in idx], "Poids actuel": cur.reindex(idx).fillna(0.0),
                       "Poids cible": tgt.reindex(idx).fillna(0.0)}, index=idx)
    df["Écart"] = df["Poids cible"] - df["Poids actuel"]
    df.index.name = "Ticker"
    return df.sort_values("Poids cible", ascending=False)


def summary_rows(res: dict) -> list[tuple[str, str]]:
    rb = res["avant"].get("risque") or {}
    ra = res["apres"].get("risque") or {}
    cons = res["avant"]["contraintes"]
    f = lambda x: "—" if x is None or (isinstance(x, float) and not np.isfinite(x)) else f"{x:.1%}"
    from utils.plain_summary import plain_lines
    simple = [("En bref" if i == 0 else "", t) for i, t in enumerate(
        x for x in plain_lines(res.get("synthese")) if not x.startswith("Lecture :"))]
    return simple + [
        ("Fichier analysé", res["fichier"]), ("Date de l'analyse", res["date"]),
        ("Profil retenu", f"{cons['profil']} — {res['source_profil']}"),
        ("Régime économique", res["avant"]["regime"]),
        ("Valeur totale", f"{res['nav']:,.0f} {res['devise']}".replace(",", " ")),
        ("Verdict du comité (actuel)", res["avant"]["verdict_label"]),
        ("Verdict du comité (après réallocation)", res["apres"]["verdict_label"]),
        ("Perte maximale acceptable (1 an)", f"{cons['perte_max']:.0%}"),
        ("Perte potentielle 1 an (VaR 95 %) actuelle / après", f"{f(rb.get('var95_1an'))} / {f(ra.get('var95_1an'))}"),
        ("Expected Shortfall 97,5 % 10 jours actuel / après", f"{f(rb.get('es975_10j'))} / {f(ra.get('es975_10j'))}"),
        ("Données de marché", "SIMULÉES — résultats non fiables" if res["data_sim"] else "réelles"),
    ]


def build_excel(res: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    NAVY, RED = PatternFill("solid", fgColor="0B1F3A"), PatternFill("solid", fgColor="FDE2E4")
    H = Font(bold=True, color="FFFFFF")

    def warn(ws, row=1, width=10):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        c = ws.cell(row=row, column=1, value="⚠ " + WARNING)
        c.fill, c.font = RED, Font(bold=True, color="9F1239")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 62

    def put(ws, df: pd.DataFrame, r0: int, c0: int, title: str, pct: tuple = (), money: tuple = (), index=True):
        ws.cell(row=r0, column=c0, value=title).font = Font(bold=True, size=12, color="0E7490")
        d = df.reset_index() if index else df
        for j, col in enumerate(d.columns):
            c = ws.cell(row=r0 + 1, column=c0 + j, value=str(col))
            c.fill, c.font = NAVY, H
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, row in enumerate(d.itertuples(index=False)):
            for j, v in enumerate(row):
                if isinstance(v, (float, np.floating)) and not np.isfinite(v):
                    v = None
                elif isinstance(v, (np.floating, np.integer)):
                    v = v.item()
                c = ws.cell(row=r0 + 2 + i, column=c0 + j, value=v)
                col = d.columns[j]
                if col in pct:
                    c.number_format = "0.0%"
                elif col in money:
                    c.number_format = "#,##0"
        return r0 + 2 + len(d)

    def widths(ws, w):
        for i, x in enumerate(w, 1):
            ws.column_dimensions[get_column_letter(i)].width = x

    # ── Synthèse : tableau du portefeuille + secteurs et régions à côté ──────
    ws = wb.active
    ws.title = "Synthèse"
    warn(ws, 1, 20)
    ws.cell(row=2, column=1, value="Conseil du comité MAM sur un portefeuille importé").font = Font(bold=True, size=15, color="0B1F3A")
    for i, (k, v) in enumerate(summary_rows(res)):
        ws.cell(row=3 + i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=3 + i, column=2, value=v)
    r = 3 + len(summary_rows(res)) + 1
    lines = res["lignes"][["Ticker", "Nom", "Secteur", "Classe", "Région", "Quantité", "Cours", "Valeur", "Poids",
                           "Plus-value", "Plus-value %"]]
    end = put(ws, lines, r, 1, "Portefeuille importé", pct=("Poids", "Plus-value %"), money=("Valeur", "Plus-value"),
              index=False)
    ws.cell(row=end, column=1, value="Liquidités").font = Font(bold=True)
    ws.cell(row=end, column=8, value=float(res["cash"])).number_format = "#,##0"
    sec = res["secteurs"][["ETF", "Perf. 3 mois", "Perf. 12 mois", "Lecture du contexte", "Score"]]
    e2 = put(ws, sec, r, 13, "Secteurs qui performent le mieux (contexte inclus)", pct=("Perf. 3 mois", "Perf. 12 mois"))
    reg = res["regions"][["Zone", "ETF", "Perf. 3 mois", "Perf. 12 mois", "Lecture du contexte", "Score"]]
    put(ws, reg, e2 + 2, 13, "Régions qui performent le mieux (contexte inclus)", pct=("Perf. 3 mois", "Perf. 12 mois"))
    widths(ws, [46, 26, 30, 22, 18, 10, 11, 13, 9, 12, 11, 3, 28, 17, 12, 12, 18, 12, 9])
    # camembert par classe
    cls = res["classes_avant"][res["classes_avant"] > 0.0005]
    ws2 = wb.create_sheet("Réallocation")
    warn(ws2, 1, 8)
    cmp_cls = pd.DataFrame({"Actuel": res["classes_avant"], "Après réallocation": res["classes_apres"]})
    cmp_cls = cmp_cls[(cmp_cls.abs().sum(axis=1) > 0.0005)]
    cmp_cls.index.name = "Classe d'actifs"
    e = put(ws2, cmp_cls, 3, 1, "Allocation par classe d'actifs", pct=("Actuel", "Après réallocation"))
    ch = BarChart(); ch.type = "bar"; ch.title = "Allocation avant / après"; ch.y_axis.number_format = "0%"
    ch.add_data(Reference(ws2, min_col=2, max_col=3, min_row=4, max_row=e - 1), titles_from_data=True)
    ch.set_categories(Reference(ws2, min_col=1, min_row=5, max_row=e - 1)); ch.height, ch.width = 9, 16
    ws2.add_chart(ch, "F3")
    pie = PieChart(); pie.title = "Portefeuille actuel par classe"
    pie.add_data(Reference(ws2, min_col=2, min_row=4, max_row=e - 1), titles_from_data=True)
    pie.set_categories(Reference(ws2, min_col=1, min_row=5, max_row=e - 1)); pie.height, pie.width = 9, 12
    ws.add_chart(pie, f"A{end + 3}")
    e = put(ws2, _realloc_frame(res), e + 2, 1, "Lignes : poids actuel et poids cible du comité",
            pct=("Poids actuel", "Poids cible", "Écart"))
    od = res["opt"]["ordres"]
    if od is not None and not od.empty:
        put(ws2, od.rename(columns={"ticker": "Ticker", "action": "Sens", "qty": "Quantité", "price": "Cours",
                                    "montant": "Montant", "motif": "Motif"}), e + 2, 1,
            "Ordres proposés (aucun n'est exécuté)", money=("Montant",), index=False)
    widths(ws2, [26, 30, 14, 14, 12, 14, 30])
    # ── Avis du comité ───────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Avis du comité")
    warn(ws3, 1, 6)
    e = put(ws3, _avis_frame(res["avant"]), 3, 1, f"Portefeuille actuel — {res['avant']['verdict_label']}", index=False)
    e = put(ws3, _avis_frame(res["apres"]), e + 2, 1, f"Après réallocation — {res['apres']['verdict_label']}", index=False)
    recos = pd.DataFrame(res["avant"]["recommandations"])
    if not recos.empty:
        put(ws3, recos[["qui", "action", "priorite"]].rename(columns={"qui": "Membre", "action": "Recommandation",
                                                                   "priorite": "Priorité"}), e + 2, 1,
            "Recommandations (portefeuille actuel)", index=False)
    widths(ws3, [24, 30, 22, 40, 90, 60])
    for row in ws3.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    # ── Scénarios ────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Scénarios")
    warn(ws4, 1, 8)
    scn = res["scenarios"].copy(); scn.index.name = "Scénario"
    e = put(ws4, scn, 3, 1, "Scénarios de marché (variation de la valeur du portefeuille)",
            pct=("Portefeuille actuel", "Après réallocation", "Écart"), money=("Perte / gain actuel ($)",))
    ch = BarChart(); ch.type = "bar"; ch.title = "Impact des scénarios"; ch.y_axis.number_format = "0%"
    ch.add_data(Reference(ws4, min_col=2, max_col=3, min_row=4, max_row=e - 1), titles_from_data=True)
    ch.set_categories(Reference(ws4, min_col=1, min_row=5, max_row=e - 1)); ch.height, ch.width = 9, 18
    ws4.add_chart(ch, "H3")
    widths(ws4, [30, 18, 18, 12, 20])
    # ── Secteurs et régions détaillés ────────────────────────────────────────
    for name, df_ in (("Secteurs", res["secteurs"]), ("Régions", res["regions"])):
        wsx = wb.create_sheet(name)
        warn(wsx, 1, 10)
        d = df_.drop(columns=[c for c in ("Contexte", "Classe", "ISO3") if c in df_.columns])
        e = put(wsx, d, 3, 1, f"{name} classés (40 % perf. 3 mois, 20 % perf. 12 mois, 40 % contexte)",
                pct=("Perf. 1 mois", "Perf. 3 mois", "Perf. 12 mois"))
        ch = BarChart(); ch.type = "bar"; ch.title = f"{name} : score (0–100)"
        sc_col = list(d.reset_index().columns).index("Score") + 1
        ch.add_data(Reference(wsx, min_col=sc_col, min_row=4, max_row=e - 1), titles_from_data=True)
        ch.set_categories(Reference(wsx, min_col=1, min_row=5, max_row=e - 1)); ch.height, ch.width = 10, 16
        wsx.add_chart(ch, "K3")
        widths(wsx, [28, 18, 10, 12, 12, 12, 18, 9])
    # ── Format et mentions ───────────────────────────────────────────────────
    ws5 = wb.create_sheet("Format & mentions")
    warn(ws5, 1, 4)
    e = put(ws5, pd.DataFrame(FORMAT_ROWS, columns=["Colonne", "Statut", "Contenu", "Exemple"]), 3, 1,
            "Format attendu du fichier importé", index=False)
    for i, rule in enumerate(FORMAT_RULES):
        ws5.cell(row=e + 1 + i, column=1, value="• " + rule)
    ws5.cell(row=e + 2 + len(FORMAT_RULES), column=1, value=CREDIT).font = Font(italic=True)
    ws5.cell(row=e + 3 + len(FORMAT_RULES), column=1, value=DISCLAIMER).font = Font(italic=True)
    widths(ws5, [16, 28, 90, 14])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(res: dict) -> bytes:
    out = io.StringIO()
    out.write(f"# {WARNING}\n# {CREDIT}\n\n")
    out.write("SYNTHÈSE\n")
    pd.DataFrame(summary_rows(res), columns=["Élément", "Valeur"]).to_csv(out, sep=";", index=False)
    blocks = [("PORTEFEUILLE IMPORTÉ", res["lignes"].drop(columns=["Base de valorisation"]), False),
              ("SECTEURS QUI PERFORMENT LE MIEUX", res["secteurs"].drop(columns=["Contexte"]), True),
              ("RÉGIONS QUI PERFORMENT LE MIEUX", res["regions"].drop(columns=["Contexte", "ISO3"]), True),
              ("RÉALLOCATION PAR LIGNE", _realloc_frame(res), True),
              ("SCÉNARIOS", res["scenarios"], True),
              ("AVIS DU COMITÉ (ACTUEL)", _avis_frame(res["avant"]), False)]
    od = res["opt"]["ordres"]
    if od is not None and not od.empty:
        blocks.insert(4, ("ORDRES PROPOSÉS (AUCUN N'EST EXÉCUTÉ)", od, False))
    for title, df, idx in blocks:
        out.write(f"\n{title}\n")
        df.to_csv(out, sep=";", index=idx, decimal=",", float_format="%.6g")
    out.write(f"\n# {WARNING}\n")
    return out.getvalue().encode("utf-8-sig")


def build_png(res: dict) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import textwrap

    BG, INK, MUT = "#0b1220", "#e2e8f0", "#94a3b8"
    from utils.plain_summary import plain_lines
    simple = [x for x in plain_lines(res.get("synthese")) if not x.startswith("Lecture :")][:6]
    wrapped = [textwrap.wrap(t, 140) for t in simple]
    n_lines = sum(len(w) for w in wrapped)
    extra = 0.25 * n_lines + (0.5 if simple else 0)
    H = 15 + extra
    fig = plt.figure(figsize=(12, H), dpi=110, facecolor=BG)
    gs = fig.add_gridspec(4, 2, height_ratios=[0.55, 1, 1, 1], hspace=0.55, wspace=0.75,
                          left=0.2, right=0.97, top=1 - (0.06 * 15 + extra) / H, bottom=0.11 * 15 / H)
    fig.text(0.04, 1 - 0.375 / H, "MAM — Conseil du comité sur un portefeuille importé", color=INK, fontsize=17,
             weight="bold")
    fig.text(0.04, 1 - 0.645 / H, f"{res['fichier']} · {res['date']} · profil {res['avant']['contraintes']['profil']} · "
                                 f"régime {res['avant']['regime']}", color=MUT, fontsize=10)
    if simple:  # En bref, en mots simples (v5.6)
        y = 1 - 1.0 / H
        fig.text(0.04, y, "EN BREF", color="#00ff88", fontsize=10.5, weight="bold", va="top")
        y -= 0.3 / H
        for i, w in enumerate(wrapped):
            for j, line in enumerate(w):
                fig.text(0.04 if j == 0 else 0.055, y, ("• " if j == 0 and i > 0 else "") + line,
                         color=INK, fontsize=9.5, weight=("bold" if i == 0 else "normal"), va="top")
                y -= 0.25 / H
    ax = fig.add_subplot(gs[0, :]); ax.axis("off")
    rb, ra = res["avant"].get("risque") or {}, res["apres"].get("risque") or {}
    clean = lambda v: re.sub(r"^[^A-Za-zÀ-ÿ]+", "", str(v)).capitalize()
    fp = lambda x: "—" if x is None or (isinstance(x, float) and not np.isfinite(x)) else f"{x:.1%}"
    cells = [("Valeur totale", f"{res['nav']:,.0f} {res['devise']}".replace(",", " ")),
             ("Profil retenu", res["avant"]["contraintes"]["profil"]),
             ("Verdict actuel", clean(res["avant"]["verdict_label"])),
             ("Verdict après réallocation", clean(res["apres"]["verdict_label"])),
             ("Perte max. acceptable (1 an)", f"{res['avant']['contraintes']['perte_max']:.0%}"),
             ("VaR 95 % 1 an : actuel → après", f"{fp(rb.get('var95_1an'))} → {fp(ra.get('var95_1an'))}")]
    for i, (k, v) in enumerate(cells):
        x0, y0 = (0.0 if i % 2 == 0 else 0.5), 0.9 - (i // 2) * 0.34
        ax.text(x0, y0, k, color=MUT, fontsize=10, transform=ax.transAxes, va="top")
        ax.text(x0 + 0.24, y0, v, color=INK, fontsize=10.5, weight="bold", transform=ax.transAxes, va="top")

    def style(a, title):
        a.set_facecolor(BG); a.set_title(title, color=INK, fontsize=11, loc="left", weight="bold")
        a.tick_params(colors=MUT, labelsize=8.5)
        for s in a.spines.values():
            s.set_color("#334155")

    a = fig.add_subplot(gs[1, :])
    cmp_cls = pd.DataFrame({"Actuel": res["classes_avant"], "Après": res["classes_apres"]})
    cmp_cls = cmp_cls[cmp_cls.abs().sum(axis=1) > 0.0005]
    y = np.arange(len(cmp_cls))
    a.barh(y - 0.2, cmp_cls["Actuel"] * 100, 0.4, color="#64748b", label="Actuel")
    a.barh(y + 0.2, cmp_cls["Après"] * 100, 0.4, color="#00d4ff", label="Après réallocation")
    a.set_yticks(y, [textwrap.shorten(c, 24, placeholder="…") for c in cmp_cls.index]); a.invert_yaxis()
    a.set_xlabel("%", color=MUT)
    a.legend(facecolor=BG, labelcolor=INK, fontsize=8.5, frameon=False)
    style(a, "Allocation par classe : actuelle et proposée par le comité")
    a = fig.add_subplot(gs[2, 0])
    scn = res["scenarios"]
    y = np.arange(len(scn))
    a.barh(y - 0.2, scn["Portefeuille actuel"] * 100, 0.4, color="#ff3b6b", label="Actuel")
    a.barh(y + 0.2, scn["Après réallocation"] * 100, 0.4, color="#00ff88", label="Après")
    a.set_yticks(y, [textwrap.shorten(s, 24, placeholder="…") for s in scn.index]); a.invert_yaxis()
    a.axvline(0, color="#475569", lw=0.8)
    a.legend(facecolor=BG, labelcolor=INK, fontsize=8, frameon=False)
    style(a, "Scénarios de marché (variation, %)")
    a = fig.add_subplot(gs[2, 1])
    lw = res["lignes"].head(10)
    a.barh(np.arange(len(lw)), lw["Poids"] * 100, color="#7c3aed")
    a.set_yticks(np.arange(len(lw)), lw["Ticker"]); a.invert_yaxis()
    style(a, "Dix premières lignes (% du portefeuille)")
    for col, key, title in ((0, "secteurs", "Secteurs porteurs (score, perf. 3 m)"),
                            (1, "regions", "Régions porteuses (score, perf. 3 m)")):
        a = fig.add_subplot(gs[3, col])
        d = res[key].head(8)
        colors = [{"favorisé": "#00ff88", "pénalisé": "#ff3b6b"}.get(v, "#ffd700") for v in d["Lecture du contexte"]]
        a.barh(np.arange(len(d)), d["Score"], color=colors)
        short = {"Consommation de base": "Conso. de base", "Consommation discrétionnaire": "Conso. discrét."}
        lab = [f"{short.get(i, i)} ({p:+.0%})" if np.isfinite(p) else short.get(i, i)
               for i, p in zip(d.index, d["Perf. 3 mois"])]
        a.set_yticks(np.arange(len(d)), lab); a.invert_yaxis(); a.set_xlim(0, 100)
        style(a, title)
    fig.text(0.04, 0.083 * 15 / H, "Vert : favorisé par le contexte · jaune : neutre · rouge : pénalisé.", color=MUT,
             fontsize=8.5)
    fig.text(0.04, 0.012 * 15 / H, "\n".join(textwrap.wrap("⚠ " + WARNING, 175)) + "\n" + CREDIT, color="#fca5a5", fontsize=8.2,
             va="bottom")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=BG)
    plt.close(fig)
    return buf.getvalue()
